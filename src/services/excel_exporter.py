from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.services.sts_database import STSDatabase

HEADERS = [
    "Sözleşme Adı / Kontrat No", "Kullanıcı", "Yİ/YD", "Sözleşme Tipi", "Faaliyetler", "Teslimat / Kabul",
    "Sözleşme İçeriği", "Sözleşmenin İmzalandığı Tarih", "T0 Başlangıç Tarihi", "T0+Ay", "Termin Tarihi", "Durum", "Kabul Tarihi", "Not",
]

def export_sts_to_excel(db: STSDatabase, output_path: Path, progress_cb=None):
    wb = Workbook(); wb.remove(wb.active)
    components = [c.get("name", "") for c in db.list_components()]
    platforms = db.list_platforms() or ["TumSozlesmeler"]
    total=max(len(platforms),1)
    for i,p in enumerate(platforms, start=1):
        ws=wb.create_sheet(title=str(p)[:31] or "Platform")
        ws.append([""]*14 + [x for c in components for x in [c, "", ""]])
        ws.append(HEADERS + [x for _ in components for x in ["Teslim Edilecek","Teslim Edilen","Kalan"]])
        for c in ws[2]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid", fgColor="1F4E78")
        for card in db.list_contracts(p):
            rows = db.connect().execute("SELECT * FROM contract_rows WHERE contract_id=? ORDER BY sort_order,id", (int(card['id']),)).fetchall()
            for rr in rows:
                row=[rr['contract_no'] or "", card.get('user',''), "", card.get('type',''), rr['activity_name'] or "", rr['delivery_acceptance_name'] or "", rr['content'] or "", rr['signed_date'] or "", rr['t0_date'] or "", rr['t0_months'] or 0, rr['termin_date'] or "", rr['status'] or "", rr['acceptance_date'] or "", rr['note'] or ""]
                cmap={x['component_name']:(x['qty_required'],x['qty_delivered'],x['qty_remaining']) for x in [dict(r) for r in db.connect().execute("SELECT * FROM contract_row_components WHERE contract_row_id=?", (rr['id'],)).fetchall()]}
                for c in components:
                    q=cmap.get(c,(0,0,0)); row.extend([q[0],q[1],q[2]])
                ws.append(row)
        if progress_cb: progress_cb(int(i*100/total), f"Excel'e aktarılıyor: {p}")
    wb.save(str(output_path))
