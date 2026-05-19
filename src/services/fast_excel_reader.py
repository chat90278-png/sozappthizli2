from __future__ import annotations

import re
import traceback
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from src.config.app_config import COMP_SHEET, TAG_HEADERS, TAG_KIND_ASSIGN, TAG_SHEET, USERS_SHEET
from src.workers.excel_workers import normalize_sheet_name, safe_sheet_name

HEADER_ROW = 4
SUBHEADER_ROW = 5
DATA_START = 6
FIRST_COMPONENT_COL = 15


class FastExcelReader:
    def __init__(self, excel_path: Path):
        self.excel_path = Path(excel_path)
        self.wb = None

    def open_workbook(self):
        if self.wb is None:
            self.wb = load_workbook(self.excel_path, read_only=True, data_only=True, keep_links=False)
        return self.wb

    def close_workbook(self):
        if self.wb is not None:
            self.wb.close(); self.wb = None

    def _iso(self, v):
        if isinstance(v, datetime): return v.date().isoformat()
        if isinstance(v, date): return v.isoformat()
        s = str(v or "").strip()
        if re.match(r"^\d{4}-\d{2}-\d{2}$", s): return s
        return "" if not s else s

    def _num(self, v):
        try: return float(v or 0)
        except Exception: return 0.0

    def platform_names(self):
        wb=self.open_workbook(); out=[]
        blocked={"anasayfa","kullanicilar","etiketler","sistem bilesenleri","platform logolari","log","config","sistemtipleri"}
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible": continue
            n=normalize_sheet_name(ws.title)
            if ws.title.startswith("_") or n.startswith("_") or n in blocked: continue
            out.append(ws.title)
        return out

    def load_users(self):
        wb=self.open_workbook();
        if USERS_SHEET not in wb.sheetnames: return []
        out=[]
        for row in wb[USERS_SHEET].iter_rows(min_row=2, max_col=6, values_only=True):
            name=str((row[0] if row else "") or "").strip()
            if name: out.append({"name":name,"role":str((row[1] if len(row)>1 else "") or ""),"active":str((row[2] if len(row)>2 else True)).lower() not in {"0","false","pasif"},"payload_json":{"raw":list(row)}})
        return out

    def load_components(self):
        wb=self.open_workbook();
        if COMP_SHEET not in wb.sheetnames: return []
        out=[]
        for row in wb[COMP_SHEET].iter_rows(min_row=2, max_col=8, values_only=True):
            name=str((row[0] if row else "") or "").strip()
            if name: out.append({"name":name,"version":str((row[1] if len(row)>1 else "") or ""),"unit":str((row[2] if len(row)>2 else "Adet") or "Adet"),"usage":self._num(row[3] if len(row)>3 else 1),"active":str((row[4] if len(row)>4 else True)).lower() not in {"0","false","pasif"},"payload_json":{"raw":list(row)}})
        return out

    def load_tags_map(self):
        wb=self.open_workbook(); defs=[]; assigns={}
        if TAG_SHEET not in wb.sheetnames: return {"tag_defs":defs,"assignments":assigns}
        for row in wb[TAG_SHEET].iter_rows(min_row=2, max_col=len(TAG_HEADERS), values_only=True):
            try:
                kind=str((row[0] if row else "") or "").upper().strip(); name=str((row[1] if len(row)>1 else "") or "").strip()
                color=str((row[2] if len(row)>2 else "#3B82F6") or "#3B82F6")
                p=safe_sheet_name(str((row[5] if len(row)>5 else "") or "").strip()); no=str((row[6] if len(row)>6 else "") or "").strip(); ctype=str((row[7] if len(row)>7 else "") or "").strip()
                if name: defs.append({"name":name,"color":color,"kind":kind or "MANUAL"})
                if kind in {"ASSIGN", TAG_KIND_ASSIGN} and p and no and name: assigns.setdefault((p,no,ctype), []).append(name)
            except Exception:
                continue
        return {"tag_defs":defs,"assignments":assigns}

    def _component_cols(self, ws):
        cols=[]; c=FIRST_COMPONENT_COL; idx=1
        while c <= ws.max_column:
            nm=str(ws.cell(HEADER_ROW, c).value or "").strip(); sub1=normalize_sheet_name(str(ws.cell(SUBHEADER_ROW,c).value or ""))
            if nm and "teslim" in sub1:
                cols.append({"name":nm,"required_col":c,"delivered_col":c+1,"remaining_col":c+2,"sort_order":idx}); idx+=1; c+=3
            else:
                c+=1
        return cols

    def _row_kind(self, e, f):
        ne, nf = normalize_sheet_name(e), normalize_sheet_name(f)
        if ne == "genel" and "ana sozlesme toplami" in nf: return "contract_total"
        if "toplami" in nf: return "system_total"
        if nf.startswith("kabul"): return "acceptance"
        return "other"

    def _rows_for_platform(self, platform):
        ws=self.open_workbook()[platform]; cols=self._component_cols(ws); out=[]
        for r in range(DATA_START, ws.max_row+1):
            vals=[ws.cell(r, c).value for c in range(1,15)]
            no=str(vals[0] or "").strip(); f=str(vals[5] or "").strip(); e=str(vals[4] or "").strip()
            if not no and not f and not e: continue
            comps=[]
            for cc in cols:
                req=self._num(ws.cell(r,cc['required_col']).value); de=self._num(ws.cell(r,cc['delivered_col']).value); rem=self._num(ws.cell(r,cc['remaining_col']).value)
                if req!=0 or de!=0 or rem!=0:
                    comps.append({"name":cc['name'],"required":req,"delivered":de,"remaining":rem,"sort_order":cc['sort_order']})
            out.append({"platform":safe_sheet_name(platform),"source_row":r,"contract_no":no,"user_name":str(vals[1] or ""),"domestic_foreign":str(vals[2] or ""),"contract_type":str(vals[3] or ""),"activity_name":str(vals[4] or ""),"delivery_acceptance_name":str(vals[5] or ""),"content":str(vals[6] or ""),"signed_date":self._iso(vals[7]),"t0_date":self._iso(vals[8]),"t0_months":int(self._num(vals[9])),"termin_date":self._iso(vals[10]),"status":str(vals[11] or ""),"acceptance_date":self._iso(vals[12]),"note":str(vals[13] or ""),"row_kind":self._row_kind(str(vals[4] or ""), str(vals[5] or "")),"components":comps})
        return out

    def build_contract_index(self, progress_cb=None):
        out=[]
        for p in self.platform_names():
            for row in self._rows_for_platform(p):
                if row['row_kind']!='contract_total' or not row['contract_no']: continue
                out.append({"platform":row['platform'],"no":row['contract_no'],"contract_no":row['contract_no'],"user":row['user_name'],"user_name":row['user_name'],"type":row['contract_type'],"contract_type":row['contract_type'],"type_display":row['contract_type'],"link":"Ana Sözleşme","status":row['status'],"completion_date":row['termin_date'],"content":row['content'],"row":row['source_row'],"start_row":row['source_row'],"is_main":True,"tags":[],"search_text":" ".join([row['platform'],row['contract_no'],row['user_name'],row['contract_type'],row['status'],row['content']]).lower()})
        return out

    def load_contract_detail(self, platform, contract_no, start_row=None): return None, [], {}

    def import_to_sts(self, db, progress_cb=None):
        report={"platforms":0,"users":0,"components":0,"tags":0,"contracts":0,"contract_rows":0,"row_components":0,"errors":[]}
        try:
            if progress_cb: progress_cb(5,"Excel açılıyor...")
            self.open_workbook(); db.init_schema()
            now=datetime.utcnow().isoformat(timespec="seconds")
            db.set_meta("imported_from_excel_path", str(self.excel_path)); db.set_meta("imported_at", now); db.set_meta("importer_version","3.0")
            st=self.excel_path.stat(); db.set_meta("source_excel_size", st.st_size); db.set_meta("source_excel_mtime", int(st.st_mtime))
            if progress_cb: progress_cb(12,"Platformlar okunuyor...")
            platforms=self.platform_names(); report['platforms']=len(platforms)
            for p in platforms: db.upsert_platform(safe_sheet_name(p))
            if progress_cb: progress_cb(20,"Kullanıcılar aktarılıyor...")
            users=self.load_users();
            if users:
                for u in users: db.upsert_user(u)
            else: db.ensure_default_user(); users=[{"name":"Sistem"}]
            report['users']=len(users)
            if progress_cb: progress_cb(28,"Bileşenler aktarılıyor...")
            comps=self.load_components(); report['components']=len(comps)
            for c in comps: db.upsert_component(c)
            if progress_cb: progress_cb(34,"Etiketler aktarılıyor...")
            t=self.load_tags_map(); report['tags']=len(t['tag_defs'])
            for tg in t['tag_defs']: db.upsert_tag(tg)
            if progress_cb: progress_cb(40,"Sözleşmeler okunuyor...")
            for i,p in enumerate(platforms, start=1):
                if progress_cb: progress_cb(40+int(i*45/max(len(platforms),1)), f"{safe_sheet_name(p)} aktarılıyor... ({i}/{len(platforms)})")
                rows=self._rows_for_platform(p)
                block=[]; current_no=""
                for rw in rows:
                    no=rw.get('contract_no','').strip()
                    if rw['row_kind']=='contract_total' and no:
                        if block:
                            h=next((x for x in block if x['row_kind']=='contract_total'), block[0])
                            cid=db.upsert_contract_from_excel_block(h, block); db.set_contract_tags(cid, t['assignments'].get((h['platform'],h['contract_no'],h['contract_type']),[])); report['contracts']+=1
                            report['contract_rows'] += len(block); report['row_components'] += sum(len(x.get('components') or []) for x in block)
                        block=[rw]; current_no=no
                    elif no and current_no and no!=current_no:
                        h=next((x for x in block if x['row_kind']=='contract_total'), block[0])
                        cid=db.upsert_contract_from_excel_block(h, block); db.set_contract_tags(cid, t['assignments'].get((h['platform'],h['contract_no'],h['contract_type']),[])); report['contracts']+=1
                        report['contract_rows'] += len(block); report['row_components'] += sum(len(x.get('components') or []) for x in block)
                        block=[rw]; current_no=no
                    elif current_no:
                        if not no: rw['contract_no']=current_no
                        block.append(rw)
                if block:
                    h=next((x for x in block if x['row_kind']=='contract_total'), block[0])
                    cid=db.upsert_contract_from_excel_block(h, block); db.set_contract_tags(cid, t['assignments'].get((h['platform'],h['contract_no'],h['contract_type']),[])); report['contracts']+=1
                    report['contract_rows'] += len(block); report['row_components'] += sum(len(x.get('components') or []) for x in block)
            if progress_cb: progress_cb(94, "Veritabanı optimize ediliyor...")
            db.add_log("excel_import_finished","database",str(db.path),"Excel dosyası STS veritabanına aktarıldı",payload=report); db.vacuum()
            if progress_cb: progress_cb(100, "Aktarım tamamlandı.")
            return report
        except Exception as exc:
            raise RuntimeError(f"Excel STS aktarımı başarısız: {exc}\n{traceback.format_exc()}")
        finally:
            self.close_workbook()
